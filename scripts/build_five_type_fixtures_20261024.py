"""Build wholly fictional five-type source fixtures. No model, service or COM.

Run with the Codex bundled Python (reportlab, python-docx, pypdf, Pillow).
The Poppler executable is explicit; no user-installed office renderer is used.
Only the requested fixture tree is written. Existing customer material is never read.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts/five-type-current-20261024/fixtures"
NATIVE = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/native"
POPLER = NATIVE / "poppler/Library/bin/pdftoppm.exe"
DISCLAIMER = "全虚构验收材料  仅供软件功能测试  非真实法律文书"
FIRM = "虚构明序律师事务所"
COURT = "虚构市样例区人民法院"


def page(title, paragraphs, doc_type, seq, *, format="pdf", heading=""):
    return dict(title=title, paragraphs=paragraphs, doc_type=doc_type,
                catalog_seq=seq, format=format, heading=heading)


CASES = [
    dict(id="civil_scan", kind="civil", title="虚构星禾货款争议一审", stage="first_instance", scan=True,
         catalog=dict(variant="civil", printed_rows=17, execution_material_present=False), pages=[
        page("委托代理通知书", [
            "致：虚构明序律师事务所。",
            "委托人：虚构星禾农业有限公司。本公司是本案被告。",
            "本公司委托贵所代理与原告虚构青禾设备有限公司之间的买卖合同纠纷一审案件。",
            "贵所指派虚构律师甲为本案承办律师。案号：（2026）粤0999民初1001号。",
            "通知日期：2026年4月2日。本通知不附委托代理合同。",
            "委托人盖章栏：空白。"], "agency_notice", 3),
        page("授权委托书", [
            "委托人：虚构星禾农业有限公司，诉讼身份为被告。",
            "受托人：虚构明序律师事务所指派的虚构律师甲。",
            "委托事项：代理本公司参加（2026）粤0999民初1001号买卖合同纠纷一审诉讼。",
            "授权范围：代为提交答辩意见、提交证据及参加庭审；特别授权事项未列明。",
            "填写日期：2026年4月3日。委托人签字盖章栏：空白。"], "poa", 4),
        page("民事起诉状", [
            "案号：（2026）粤0999民初1001号。",
            "原告：虚构青禾设备有限公司。被告：虚构星禾农业有限公司。",
            "案由：买卖合同纠纷。诉讼请求：请求法院判令被告支付货款30000元。",
            "原告诉称其已交付设备而被告尚欠货款。上述内容为原告主张。",
            "原告表示：如取得生效裁判而被告未履行，将另行申请执行。",
            "落款日期：2026年3月28日。"], "complaint", 5),
        page("民事判决书", [
            "（2026）粤0999民初1001号",
            "原告：虚构青禾设备有限公司。被告：虚构星禾农业有限公司。",
            "案由：买卖合同纠纷。被告委托诉讼代理人：虚构律师甲，虚构明序律师事务所律师。",
            "本院判决如下：被告虚构星禾农业有限公司于本判决生效之日起十日内向原告虚构青禾设备有限公司支付货款20000元；驳回原告其他诉讼请求。",
            "裁判日期：2026年9月30日。",
            "本页没有生效日期证明、付款回单或执行立案材料。"], "judgment", 14, heading=COURT),
    ]),
    dict(id="admin_mixed", kind="admin", title="虚构远山行政处罚及执行", stage="execution", scan=False,
         catalog=dict(variant="admin_execution", printed_rows=18, execution_material_present=True), pages=[
        page("委托代理合同", [
            "委托人：虚构远山包装有限公司。受托人：虚构明序律师事务所。",
            "委托事项：本公司作为原告提起行政处罚撤销及返还罚款诉讼，并办理相应执行事项。",
            "行政相对方：虚构市样例区市场监督管理局。行政案由：行政处罚争议。",
            "本所指派虚构律师乙为本案承办律师。约定律师费6000元，收款凭证未附。",
            "合同填写日期：2026年5月10日。",
            "甲方签字盖章栏：空白。乙方签字盖章栏：空白。"], "contract", 3, format="docx"),
        page("行政起诉状", [
            "（2026）粤0999行初2001号",
            "原告：虚构远山包装有限公司。被告：虚构市样例区市场监督管理局。",
            "行政案由：行政处罚争议。",
            "诉讼请求：撤销虚市监罚字〔2026〕第010号处罚决定，并请求返还罚款9000元。",
            "原告主张涉案处罚事实认定有误。以上为原告请求，尚非裁判结论。",
            "起诉日期：2026年5月12日。"], "complaint", 5),
        page("行政判决书", [
            "（2026）粤0999行初2001号",
            "原告：虚构远山包装有限公司。被告：虚构市样例区市场监督管理局。",
            "本院判决如下：撤销虚市监罚字〔2026〕第010号行政处罚决定；被告于本判决生效之日起三十日内向原告返还罚款9000元。",
            "裁判日期：2026年7月20日。",
            "本页没有记载实际返还到账日期。"], "judgment", 14, format="png", heading=COURT),
        page("执行通知书", [
            "（2026）粤0999执3001号",
            "申请执行人：虚构远山包装有限公司。被执行人：虚构市样例区市场监督管理局。",
            "执行依据：（2026）粤0999行初2001号行政判决书。",
            "本院已立案执行申请执行人与被执行人之间返还罚款9000元一案，责令被执行人履行生效法律文书确定的义务。",
            "通知日期：2026年9月8日。本通知不记载款项到账或执行完毕。"], "execution", 15, heading=COURT),
    ]),
    dict(id="criminal_mixed", kind="criminal", title="虚构程乙刑事辩护一审", stage="first_instance", scan=False,
         catalog=dict(variant="criminal", printed_rows=18, execution_material_present=False), pages=[
        page("刑事辩护委托合同", [
            "委托人暨本案被告人：虚构程乙。受托人：虚构明序律师事务所。",
            "委托事项：就虚构程乙被控故意毁坏财物罪一案提供一审辩护。",
            "本所接受虚构程乙的委托，指派虚构律师丙担任本案辩护人。",
            "法院案号：（2026）粤0999刑初4001号。约定辩护费5000元，未附收款凭证。",
            "合同填写日期：2026年6月2日。双方签字盖章栏：空白。"], "contract", 3),
        page("起诉书", [
            "虚检刑诉〔2026〕第040号",
            "被告人：虚构程乙。法院案号：（2026）粤0999刑初4001号。",
            "指控罪名：故意毁坏财物罪。",
            "本院指控虚构程乙毁损他人设备，请求法院依法判处相应刑罚。该刑罚请求不等于法院已经判决。",
            "提起公诉日期：2026年6月15日。",
            "出具机关：虚构市样例区人民检察院。"], "indictment", 5, format="docx"),
        page("会见笔录", [
            "会见时间：2026年6月20日。会见地点：虚构市样例看守所。",
            "会见人：虚构明序律师事务所受虚构程乙委托指派的辩护人虚构律师丙。",
            "被会见人：虚构程乙；本案被告人。案号：（2026）粤0999刑初4001号。",
            "会见内容：虚构程乙陈述设备损坏经过，并表示希望争取从轻处罚；该陈述未经本页独立证明。",
            "被会见人确认签字栏：空白。"], "meeting_record", 7, format="png"),
        page("辩护词", [
            "（2026）粤0999刑初4001号",
            "虚构明序律师事务所接受被告人虚构程乙委托，指派虚构律师丙担任本案一审辩护人。",
            "被控罪名：故意毁坏财物罪。",
            "辩护意见：请求法院结合被告人陈述及在案证据，对虚构程乙从轻处罚并适用缓刑。此为辩护请求。",
            "提交日期：2026年8月12日。"], "plea", 12),
        page("刑事判决书", [
            "（2026）粤0999刑初4001号",
            "公诉机关：虚构市样例区人民检察院。被告人：虚构程乙。",
            "辩护人：虚构律师丙，虚构明序律师事务所律师，受被告人虚构程乙委托。",
            "本院判决如下：被告人虚构程乙犯故意毁坏财物罪，判处拘役三个月，缓刑六个月。",
            "裁判日期：2026年9月10日。",
            "本页未载明判决生效日期、刑罚执行结束或损害赔偿到账事实。"], "judgment", 14, heading=COURT),
    ]),
    dict(id="nonlit_mixed", kind="nonlit", title="虚构星桥采购合同审查", stage="delivery", scan=False,
         catalog=dict(variant="nonlit", printed_rows=10, execution_material_present=False), pages=[
        page("非诉讼法律事务委托合同", [
            "委托人：虚构星桥商贸有限公司。受托人：虚构明序律师事务所。",
            "项目事项：采购合同审查。项目名称：虚构星桥采购合同审查项目。",
            "本所接受该公司委托，指派虚构律师丁办理本项目。",
            "工作范围：审查采购合同草案并出具法律意见及修改稿；不包括代理诉讼或办理付款。",
            "约定律师费2000元，未附收款凭证。合同填写日期：2026年9月1日。",
            "双方签字盖章栏：空白。"], "contract", 2, format="docx"),
        page("与委托人谈话记录", [
            "谈话日期：2026年9月3日。项目：虚构星桥采购合同审查项目。",
            "参加人：虚构星桥商贸有限公司项目联系人虚构联络甲；本所承办律师虚构律师丁。",
            "委托人提出希望在采购合同中明确验收标准与违约责任。",
            "委托人说明当前采购合同仍为草案，尚未提供双方签署版本。",
            "记录人：虚构律师丁。委托人确认签字栏：空白。"], "client_talk", 4, format="png"),
        page("法律意见书", [
            "致：虚构星桥商贸有限公司。项目事项：采购合同审查。",
            "虚构明序律师事务所受贵公司委托，指派虚构律师丁审查采购合同草案。",
            "审查结论：现有草案的验收期限和违约责任约定不够明确，建议在签署前补充具体条款。",
            "本意见针对送审草案，不证明采购交易已经签署、货物已经交付或货款已经支付。",
            "出具日期：2026年9月10日。"], "legal_work", 7),
        page("办理法律事务活动记录", [
            "项目名称：虚构星桥采购合同审查项目。委托人：虚构星桥商贸有限公司。",
            "承办律师：虚构明序律师事务所接受本项目委托后指派的虚构律师丁。",
            "办理结果：2026年9月12日完成采购合同草案审查，并向委托人发出法律意见书和合同修改稿。",
            "本记录未附委托人签收回执；采购合同最终签订情况及律师费收款情况未记载。",
            "记录日期：2026年9月12日。"], "legal_work", 7),
    ]),
    dict(id="counsel_mixed", kind="counsel", title="虚构星岚年度法律顾问", stage="service_period", scan=False,
         catalog=dict(variant="counsel", printed_rows=9, execution_material_present=False), pages=[
        page("法律顾问合同", [
            "聘方暨委托人：虚构星岚科技有限公司。受聘方：虚构明序律师事务所。",
            "顾问事项：日常合同审查与经营法律咨询。",
            "本所接受虚构星岚科技有限公司聘请，指派虚构律师戊担任本单位法律顾问律师。",
            "服务期限：2026年1月1日至2026年12月31日。约定顾问费12000元，未附支付凭证。",
            "合同填写日期：2025年12月20日。双方签字盖章栏：空白。"], "contract", 2),
        page("聘方基本情况介绍材料", [
            "聘方名称：虚构星岚科技有限公司。",
            "本材料由聘方向虚构明序律师事务所提供，用于年度法律顾问工作。",
            "联系地址：虚构市样例区演示路88号。联系电话：未提供。",
            "经营情况介绍：从事演示设备研发；本段为虚构聘方提供的说明，未附登记证照。",
            "顾问事项：日常合同审查与经营法律咨询。提供日期：2026年1月5日。"], "client_intro", 3, format="docx"),
        page("法律顾问工作记录", [
            "服务对象：虚构星岚科技有限公司。顾问事项：日常合同审查与经营法律咨询。",
            "虚构明序律师事务所接受聘请，指派虚构律师戊作为本单位顾问律师。",
            "2026年10月10日，虚构律师戊完成采购合同模板审查，向聘方发送修订意见。",
            "本次工作只涉及模板审查；不涉及诉讼案件，未附聘方签收证明。",
            "记录日期：2026年10月10日。"], "work_record", 5, format="png"),
        page("法律顾问协议存续情况说明", [
            "聘方：虚构星岚科技有限公司。受聘方：虚构明序律师事务所。",
            "对应合同：服务期限为2026年1月1日至2026年12月31日的法律顾问合同。",
            "截至2026年10月20日，该顾问协议仍在履行，尚未终止；本说明不是终止协议。",
            "已完成的阶段工作包括2026年10月10日的采购合同模板审查。后续日常顾问服务继续。",
            "说明日期：2026年10月20日。顾问费实际支付情况未记载。"], "agreement_status", 6),
    ]),
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def wrapped(text, font="FixtureSong", size=14, width=475):
    line = ""
    for char in text:
        if pdfmetrics.stringWidth(line + char, font, size) > width:
            if char in "。，；：！？、）】》":
                yield line + char
                line = ""
                continue
            yield line
            line = char
        else:
            line += char
    if line:
        yield line


def draw_pdf(path, item, source_label):
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle(item["title"] + " 全虚构验收材料")
    c.setAuthor("GD synthetic fixtures")
    y = 790
    if item["heading"]:
        c.setFont("FixtureHei", 16)
        c.drawCentredString(A4[0] / 2, y, item["heading"])
        y -= 32
    c.setFont("FixtureHei", 21)
    c.drawCentredString(A4[0] / 2, y, item["title"])
    y -= 34
    c.setFont("FixtureSong", 9)
    c.drawCentredString(A4[0] / 2, y, DISCLAIMER)
    y -= 40
    c.setFont("FixtureSong", 14)
    for para in item["paragraphs"]:
        for line in wrapped(para):
            c.drawString(60, y, line)
            y -= 25
        y -= 13
    if y < 85:
        raise AssertionError(f"Fixture overflows: {source_label}: {y}")
    c.setFont("FixtureSong", 9)
    c.drawString(60, 42, "全虚构  无真实签章  " + source_label)
    c.save()


def render(pdf, stem, dpi=120):
    subprocess.run([str(POPLER), "-f", "1", "-singlefile", "-r", str(dpi), "-png", str(pdf), str(stem)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    result = stem.with_suffix(".png")
    if not result.is_file():
        raise AssertionError("Poppler produced no image")
    return result


def docx(path, item):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    for style in [d.styles["Normal"], d.styles["Title"]]:
        style.font.name = "SimSun"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    d.styles["Normal"].font.size = Pt(12)
    d.styles["Normal"].paragraph_format.line_spacing = 1.4
    d.styles["Normal"].paragraph_format.space_after = Pt(10)
    d.styles["Title"].font.size = Pt(20)
    d.add_paragraph(item["title"], "Title")
    mark = d.add_paragraph(DISCLAIMER)
    mark.runs[0].font.size = Pt(9)
    for para in item["paragraphs"]:
        d.add_paragraph(para)
    d.core_properties.author = "GD synthetic fixtures"
    d.core_properties.title = item["title"] + " 全虚构验收材料"
    d.core_properties.created = d.core_properties.modified = datetime(2026, 10, 24, tzinfo=timezone.utc)
    raw = io.BytesIO()
    d.save(raw)
    # Stable ZIP metadata, no user identity or local source path.
    with ZipFile(raw) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, (2026, 10, 24, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, source.read(name))
    assert [p.text for p in Document(path).paragraphs] == [item["title"], DISCLAIMER, *item["paragraphs"]]


def expected_for(case, docs):
    def source(index):
        return {"file": docs[index]["file"], "page": docs[index]["page_start"]}

    def field(name, value, index, quote, *, match="contains", aliases=None):
        return dict(field=name, value=value, source=source(index), evidence_quote=quote,
                    match=match, aliases=aliases or [])

    def forbidden(fact, reason, indexes):
        return dict(fact=fact, reason=reason, sources=[source(i) for i in indexes])

    common = [
        forbidden("委托合同已签署或授权书已签字盖章", "合同/授权材料的存在不证明签署；样例签字盖章栏为空白。", [0]),
        forbidden("律师费已收取或清偿", "约定费用或委托通知不是收款凭证。", [0]),
        forbidden("收案日期等于合同填写日或通知日期", "未提供独立收案记录。", [0]),
        forbidden("系统生成质量监督卡或送达材料清单证明签署、签收或送达", "原始材料未提供相应签署/签收凭证；生成空表不证明历史行为。", [0]),
    ]
    k = case["kind"]
    if k == "civil":
        fields = [field("委托人名称", "虚构星禾农业有限公司", 0, "委托人：虚构星禾农业有限公司。本公司是本案被告。", aliases=["委托人"]),
                  field("承办律师", "虚构律师甲", 0, "贵所指派虚构律师甲为本案承办律师。"),
                  field("案由", "买卖合同纠纷", 2, "案由：买卖合同纠纷。"),
                  field("案号", "（2026）粤0999民初1001号", 3, "（2026）粤0999民初1001号", match="normalized", aliases=["法院收案号"]),
                  field("法院判决", "20000元", 3, "本院判决如下：被告虚构星禾农业有限公司于本判决生效之日起十日内向原告虚构青禾设备有限公司支付货款20000元；驳回原告其他诉讼请求。", aliases=["结案小结"])]
        forbidden_values = common + [forbidden("本所代理原告或委托人是虚构青禾设备有限公司", "原页明确本所受被告委托。", [0, 1]),
            forbidden("已申请执行、已执行完毕或已付款20000元", "只有起诉状中条件性执行表述和判令支付，未附执行/付款凭证。", [2, 3]),
            forbidden("法院判令支付30000元", "30000元是原告诉请；裁判金额是20000元。", [2, 3]),
            forbidden("委托代理合同已提供", "第1页是委托代理通知书；可归合同目录槽3，但本包没有合同。", [0])]
    elif k == "admin":
        fields = [field("委托人名称", "虚构远山包装有限公司", 0, "委托人：虚构远山包装有限公司。受托人：虚构明序律师事务所。", aliases=["委托人"]),
                  field("承办律师", "虚构律师乙", 0, "本所指派虚构律师乙为本案承办律师。"),
                  field("案由", "行政处罚争议", 0, "行政案由：行政处罚争议。"),
                  field("案号", "（2026）粤0999行初2001号", 2, "（2026）粤0999行初2001号", match="normalized", aliases=["法院收案号"]),
                  field("执行案号", "（2026）粤0999执3001号", 3, "（2026）粤0999执3001号", match="normalized")]
        forbidden_values = common + [forbidden("实际收回9000元、行政机关已退款或执行完毕", "执行通知只证明存在本案执行材料，未证明款项到账。", [2, 3]),
            forbidden("以执行案号替换行政审判案号", "两份文书载明不同程序的案号，应区分来源与阶段。", [2, 3])]
    elif k == "criminal":
        fields = [field("委托人名称", "虚构程乙", 0, "委托人暨本案被告人：虚构程乙。", aliases=["委托人"]),
                  field("承办律师", "虚构律师丙", 0, "本所接受虚构程乙的委托，指派虚构律师丙担任本案辩护人。"),
                  field("案由", "故意毁坏财物罪", 1, "指控罪名：故意毁坏财物罪。", aliases=["罪名"]),
                  field("案号", "（2026）粤0999刑初4001号", 4, "（2026）粤0999刑初4001号", match="normalized", aliases=["法院收案号"]),
                  field("法院判决", "拘役三个月，缓刑六个月", 4, "本院判决如下：被告人虚构程乙犯故意毁坏财物罪，判处拘役三个月，缓刑六个月。", aliases=["结案小结"])]
        forbidden_values = common + [forbidden("刑罚执行结束、损害赔偿到账或判决已生效", "未附相关事实证明。", [4]),
            forbidden("起诉或辩护请求直接等于法院裁判", "起诉书/辩护词为各方主张；判决另载明确主文。", [1, 3, 4])]
    elif k == "nonlit":
        fields = [field("委托人名称", "虚构星桥商贸有限公司", 0, "委托人：虚构星桥商贸有限公司。", aliases=["委托人"]),
                  field("承办律师", "虚构律师丁", 0, "本所接受该公司委托，指派虚构律师丁办理本项目。"),
                  field("案由", "采购合同审查", 0, "项目事项：采购合同审查。", aliases=["项目事项"]),
                  field("案件或项目名称", "虚构星桥采购合同审查项目", 0, "项目名称：虚构星桥采购合同审查项目。"),
                  field("结案小结", "完成采购合同草案审查", 3, "办理结果：2026年9月12日完成采购合同草案审查，并向委托人发出法律意见书和合同修改稿。")]
        forbidden_values = common + [forbidden("法院案号、诉讼原告被告或胜诉结果", "这是非诉审查项目，没有法院案件或裁判。", [0, 2, 3]),
            forbidden("采购交易已签署、履行或委托人已签收修改稿", "仅记录审查及发送，未附采购签约、履约或签收凭证。", [1, 2, 3])]
    else:
        fields = [field("委托人名称", "虚构星岚科技有限公司", 0, "聘方暨委托人：虚构星岚科技有限公司。", aliases=["委托人"]),
                  field("承办律师", "虚构律师戊", 0, "本所接受虚构星岚科技有限公司聘请，指派虚构律师戊担任本单位法律顾问律师。", aliases=["顾问律师"]),
                  field("案由", "日常合同审查与经营法律咨询", 0, "顾问事项：日常合同审查与经营法律咨询。", aliases=["顾问事项"]),
                  field("委托人联系地址及电话", "虚构市样例区演示路88号", 1, "联系地址：虚构市样例区演示路88号。联系电话：未提供。", aliases=["委托人联系地址"])]
        forbidden_values = common + [forbidden("顾问协议已终止或全部服务已办结", "2026年10月20日协议仍在履行，只完成阶段审查工作。", [2, 3]),
            forbidden("虚构法院案号、诉讼对方或联系电话", "没有诉讼案件，电话明确未提供。", [1, 2])]
    contract = k != "civil"
    return dict(catalog=case["catalog"], documents=docs, field_checks=fields,
                forbidden_inferences=forbidden_values,
                material_checks=[dict(doc_type="contract", present=contract,
                                      signature_verified=False,
                                      reason="有合同正文但签字盖章栏空白。" if contract else "仅有通知书；合同正文未提供。",
                                      source=source(0))],
                boundaries=["短材料功能链样例，不是完整业务档案或法定齐备材料。",
                            "预期是人工核对基准，不是模型实际识别结果。",
                            "没有真实身份、客户材料、印章、签名、证件号或可拨通联系电话。",
                            "不能因目录有某项、材料已分类或生成系统空表就认定材料齐备/已签署/已送达。"])


def build():
    if not POPLER.is_file():
        raise FileNotFoundError("Bundled Poppler is required")
    OUT.mkdir(parents=True, exist_ok=True)
    qa = OUT / "_qa"
    qa.mkdir(exist_ok=True)
    pdfmetrics.registerFont(TTFont("FixtureSong", "C:/Windows/Fonts/simsun.ttc", subfontIndex=0))
    pdfmetrics.registerFont(TTFont("FixtureHei", "C:/Windows/Fonts/simhei.ttf"))
    manifest = dict(schema_version=1, synthetic=True, created="2026-10-24", cases=[])
    readability = dict(source="entirely_fictional_literals_in_builder", model_calls=0, office_com_calls=0,
                       pdf_renderer="bundled Poppler", images=[], structural_checks=[],
                       native_docx_render="not_performed_no_bundled_libreoffice_on_windows",
                       native_docx_visual_status="pending_parent_actual_import_conversion",
                       visual_review="pending")
    for case in CASES:
        folder = OUT / case["id"]
        folder.mkdir(exist_ok=True)
        refs, inputs, documents = [], [], []
        for index, item in enumerate(case["pages"], 1):
            name = f"{index:02}_{item['title']}"
            refpdf = qa / f"{case['id']}_{index:02}.pdf"
            draw_pdf(refpdf, item, f"{case['id']} / {index}")
            png = render(refpdf, qa / f"{case['id']}_{index:02}")
            refs.append(png)
            readability["images"].append(dict(path=png.relative_to(OUT).as_posix(), kind="same_source_layout_reference" if item["format"] == "docx" else "actual_pdf_or_source_image", title=item["title"]))
            if case["scan"]:
                rel = f"{case['id']}/01_整卷扫描.pdf"
                local_page = index
            else:
                path = folder / f"{name}.{item['format']}"
                if item["format"] == "docx":
                    docx(path, item)
                elif item["format"] == "png":
                    path.write_bytes(png.read_bytes())
                else:
                    path.write_bytes(refpdf.read_bytes())
                rel, local_page = path.relative_to(OUT).as_posix(), 1
                inputs.append(dict(relative_path=rel, name=path.name, sha256=sha(path), pages=1,
                                   format=item["format"], scan_only=False,
                                   page_count_verified=item["format"] != "docx"))
                readability["structural_checks"].append(dict(file=rel, text_roundtrip=item["format"] == "docx", expected_pages=1,
                                                             actual_pdf_pages=len(PdfReader(path).pages) if item["format"] == "pdf" else None))
            documents.append(dict(file=rel, page_start=local_page, page_end=local_page,
                                  title=item["title"], doc_type=item["doc_type"], catalog_seq=item["catalog_seq"],
                                  evidence_quote=" ".join(item["paragraphs"])))
        if case["scan"]:
            path = folder / "01_整卷扫描.pdf"
            c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
            c.setTitle(case["title"] + " 全虚构扫描卷")
            c.setAuthor("GD synthetic fixtures")
            for png in refs:
                c.drawImage(ImageReader(str(png)), 0, 0, width=A4[0], height=A4[1])
                c.showPage()
            c.save()
            reader = PdfReader(path)
            assert len(reader.pages) == 4 and all(not p.extract_text().strip() for p in reader.pages)
            subprocess.run([str(POPLER), "-r", "120", "-png", str(path), str(qa / "civil_scan_actual")],
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
            for index in range(1, 5):
                assert (qa / f"civil_scan_actual-{index}.png").is_file()
            readability["scan_pdf_final_render"] = [f"_qa/civil_scan_actual-{index}.png" for index in range(1, 5)]
            inputs.append(dict(relative_path=path.relative_to(OUT).as_posix(), name=path.name, sha256=sha(path),
                               pages=4, format="pdf", scan_only=True, page_count_verified=True))
            readability["structural_checks"].append(dict(file=inputs[0]["relative_path"], actual_pdf_pages=4, text_layer_empty=True))
        run = dict(schema_version=1, id=case["id"], kind=case["kind"], title=case["title"], stage=case["stage"], synthetic=True,
                   input_root="fixtures", input_pages=sum(x["pages"] for x in inputs), inputs=inputs,
                   expected=expected_for(case, documents))
        dump(folder / "run.json", run)
        manifest["cases"].append(dict(id=case["id"], kind=case["kind"], title=case["title"], stage=case["stage"],
                                      run_json=f"{case['id']}/run.json", run_sha256=sha(folder / "run.json")))
    manifest["builder_sha256"] = sha(Path(__file__))
    manifest["counts"] = dict(cases=5, input_files=sum(1 if c["scan"] else len(c["pages"]) for c in CASES), source_pages=sum(len(c["pages"]) for c in CASES), docx_native_pages_pending=4)
    dump(OUT / "manifest.json", manifest)
    dump(qa / "readability.json", readability)
    (OUT / "README.txt").write_text("五类全虚构功能验收材料。仅上传各 run.json 的 inputs，不上传 run.json、manifest 或 _qa。\n"
        "民事为4页纯扫描PDF；行政4页、刑事5页、非诉4页、顾问4页均混合PDF/DOCX/PNG。\n"
        "预期清单是人工基准，不是实际模型结论，也不代表完整业务档案或材料法定齐全。\n"
        "四个DOCX已作文本/结构检查，原生分页与目视留给主任务实际导入；_qa同源图不是DOCX原生渲染。\n", encoding="utf-8")
    print(json.dumps({"manifest": str(OUT / "manifest.json"), "sha256": sha(OUT / "manifest.json"), **manifest["counts"]}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    build()
