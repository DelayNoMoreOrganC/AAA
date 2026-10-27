"""Synthetic inputs, not extraction snapshots. Every demo goes through normal import."""
from pathlib import Path
from io import BytesIO

PAGES = [
    ("委托代理合同", ["本页全部人物、机构、案号和事实均为虚构，仅供功能体验。", "合同编号：演示委托2026-001", "委托人：星河演示贸易有限公司", "律师事务所：岭南演示律师事务所", "承办律师：陈示例", "对方当事人：远山演示设备有限公司", "案由：买卖合同纠纷", "代理阶段：民事一审", "收费标准：固定律师费人民币5000元", "收案日期：2026年3月1日", "委托事项：代理委托人与对方当事人的买卖合同纠纷。", "本合同仅记载委托安排，不证明款项已支付或收到。"]),
    ("授权委托书", ["本页全部人物、机构、案号和事实均为虚构，仅供功能体验。", "委托人：星河演示贸易有限公司", "受托人：陈示例，岭南演示律师事务所律师", "对方当事人：远山演示设备有限公司", "案由：买卖合同纠纷", "授权范围：代为提交材料、参加庭审。", "授权日期：2026年3月2日"]),
    ("民事起诉状", ["本页全部人物、机构、案号和事实均为虚构，仅供功能体验。", "原告：星河演示贸易有限公司", "被告：远山演示设备有限公司", "委托代理人：陈示例，岭南演示律师事务所律师", "案由：买卖合同纠纷", "诉讼请求：判令被告支付货款人民币100000元。", "事实与理由：双方订立设备买卖合同，原告主张被告尚欠货款。", "受理法院：演示市示例区人民法院", "起诉日期：2026年3月3日"]),
    ("民事判决书", ["演示市示例区人民法院", "本页全部人物、机构、案号和事实均为虚构，仅供功能体验。", "案号：（2026）粤0000民初1001号", "原告：星河演示贸易有限公司", "被告：远山演示设备有限公司", "案由：买卖合同纠纷", "承办律师：陈示例", "判决如下：被告于本判决生效之日起十日内支付原告货款100000元。", "本文不证明判决已生效，不证明被告已经履行，也不证明原告已收款。", "裁判日期：2026年6月18日"]),
]


def page_png(title, lines):
    import fitz
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_textbox(fitz.Rect(45, 50, 550, 105), title, fontname="china-s", fontsize=22, align=1)
    page.insert_textbox(fitz.Rect(50, 130, 545, 770), "\n\n".join(lines), fontname="china-s", fontsize=12, lineheight=1.5)
    data = page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")
    document.close()
    return data


def create_demo_inputs(directory: Path, kind="mixed") -> list[Path]:
    import fitz
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    directory.mkdir(parents=True, exist_ok=True)
    images = [page_png(*page) for page in PAGES]
    if kind == "scan":
        target = directory / "虚构案件_整份扫描卷.pdf"
        pdf = fitz.open()
        for data in images:
            page = pdf.new_page(width=595, height=842)
            page.insert_image(page.rect, stream=data)
        pdf.save(target)
        pdf.close()
        return [target]
    outputs = []
    for index in range(2):
        target = directory / (PAGES[index][0] + ".pdf")
        pdf = fitz.open()
        page = pdf.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=images[index])
        pdf.save(target)
        pdf.close()
        outputs.append(target)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)
    doc.add_heading(PAGES[2][0], 0)
    for line in PAGES[2][1]:
        doc.add_paragraph(line)
    target = directory / "民事起诉状.docx"
    doc.save(target)
    outputs.append(target)
    target = directory / "民事判决书.png"
    target.write_bytes(images[3])
    outputs.append(target)
    return outputs
