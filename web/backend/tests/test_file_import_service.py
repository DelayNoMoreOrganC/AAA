"""Local-only import checks using generated PDFs, images and DOCX fixtures."""
import json
from pathlib import Path
import subprocess
import sys
import zipfile

import fitz
from PIL import Image
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services import file_import_service as imports


def pdf_bytes(text="Fictional archive"):
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), text)
        return doc.tobytes()


def make_docx(path, extra=None):
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        for name, data in (extra or {}).items():
            archive.writestr(name, data)


def test_pdf_preserves_original_bytes_and_records_provenance(tmp_path):
    original = tmp_path / "original.pdf"
    contents = pdf_bytes()
    original.write_bytes(contents)
    target = imports.normalize_import(original, "虚构卷宗.pdf")
    assert target == original
    assert target.read_bytes() == contents
    metadata = json.loads((tmp_path / "import.json").read_text(encoding="utf-8"))
    assert metadata == {"filename": "虚构卷宗.pdf", "original": "original.pdf", "pdf": "original.pdf", "pages": 1, "conversion": "none"}


def test_rejects_invalid_and_password_protected_pdfs(tmp_path):
    original = tmp_path / "original.pdf"
    original.write_bytes(b"<html>not a PDF</html>")
    with pytest.raises(imports.ImportFailure, match="有效的 PDF"):
        imports.normalize_import(original, "bad.pdf")
    with fitz.open() as document:
        document.new_page()
        document.save(str(original), encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="reader")
    with pytest.raises(imports.ImportFailure, match="加密"):
        imports.normalize_import(original, "locked.pdf")


def test_transparent_png_becomes_a4_pdf_without_changing_original(tmp_path):
    original = tmp_path / "original.png"
    with Image.new("RGBA", (80, 120), (255, 0, 0, 80)) as image:
        image.save(original)
    contents = original.read_bytes()
    target = imports.normalize_import(original, "证据照片.png")
    assert original.read_bytes() == contents
    with fitz.open(target) as document:
        assert document.page_count == 1
        assert document[0].rect.width == 595
        assert document[0].rect.height == 842
        assert document[0].get_images()


def test_tiff_keeps_all_pages_and_limits_image_size(tmp_path, monkeypatch):
    original = tmp_path / "original.tiff"
    with Image.new("RGB", (20, 30), "white") as first, Image.new("RGB", (30, 20), "black") as second:
        first.save(original, save_all=True, append_images=[second])
    target = imports.normalize_import(original, "两页证据.tiff")
    with fitz.open(target) as document:
        assert document.page_count == 2
    monkeypatch.setattr(imports, "MAX_IMAGE_PIXELS", 10)
    with pytest.raises(imports.ImportFailure, match="分辨率"):
        imports.normalize_import(original, "large.tiff")


def test_no_converter_means_docx_not_advertised_or_silently_accepted(monkeypatch):
    monkeypatch.setattr(imports, "docx_converter", lambda: (None, ""))
    capabilities = imports.import_capabilities(200)
    assert capabilities["docx_available"] is False
    assert ".docx" not in capabilities["supported_extensions"]
    with pytest.raises(imports.ImportFailure, match="未检测到") as error:
        imports.check_extension("document.docx")
    assert error.value.status_code == 503


def test_docx_rejects_external_resource_and_embedded_code(tmp_path):
    original = tmp_path / "original.docx"
    relations = '<Relationships><Relationship TargetMode="External" Type="type/image" Target="https://invalid.test/image.png"/></Relationships>'
    make_docx(original, {"word/_rels/document.xml.rels": relations})
    with pytest.raises(imports.ImportFailure, match="外部链接"):
        imports._validate_docx(original)
    make_docx(original, {"word/vbaProject.bin": b"not-real-code"})
    with pytest.raises(imports.ImportFailure, match="宏"):
        imports._validate_docx(original)


def test_docx_rejects_dynamic_external_fields_before_office_opens_it(tmp_path):
    original = tmp_path / "original.docx"
    with zipfile.ZipFile(original, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", '<document><instrText>DDEAUTO </instrText><instrText>fictional target</instrText></document>')
    with pytest.raises(imports.ImportFailure, match="动态链接域"):
        imports._validate_docx(original)


def test_docx_converter_output_must_be_a_real_pdf(tmp_path, monkeypatch):
    original = tmp_path / "original.docx"
    make_docx(original)
    monkeypatch.setattr(imports, "docx_converter", lambda: ("libreoffice", "fictional-soffice"))

    def fake_converter(command, **kwargs):
        output = Path(command[command.index("--outdir") + 1]) / "original.pdf"
        output.write_bytes(pdf_bytes())
        assert "--headless" in command
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in command)
        assert kwargs["timeout"] == 120
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(imports.subprocess, "run", fake_converter)
    target = imports.normalize_import(original, "委托书.docx")
    assert imports.validate_pdf(target) == 1
    assert original.exists()


def test_docx_converter_failure_is_not_success(tmp_path, monkeypatch):
    original = tmp_path / "original.docx"
    make_docx(original)
    monkeypatch.setattr(imports, "docx_converter", lambda: ("libreoffice", "fictional-soffice"))
    monkeypatch.setattr(imports.subprocess, "run", lambda command, **kwargs: subprocess.CompletedProcess(command, 1))
    with pytest.raises(imports.ImportFailure, match="转换失败") as error:
        imports.normalize_import(original, "failed.docx")
    assert error.value.status_code == 422
    assert original.exists()
    assert not (tmp_path / "import.json").exists()


def test_original_lookup_cannot_escape_owned_import_directory(tmp_path):
    owned = tmp_path / "imports" / ("a" * 32)
    owned.mkdir(parents=True)
    outside = tmp_path / "private.txt"
    outside.write_text("fictional")
    (owned / "import.json").write_text(json.dumps({"original": "../../private.txt"}))
    with pytest.raises(FileNotFoundError):
        imports.original_file(tmp_path, owned / "source.pdf")
    assert imports.owned_import_directory(tmp_path, tmp_path / "anything" / "source.pdf") is None


def test_display_filename_handles_windows_and_posix_paths():
    assert imports.display_filename(r"C:\fakepath\证据.pdf") == "证据.pdf"
    assert imports.display_filename("../../证据.pdf") == "证据.pdf"
    assert imports.display_filename("\x00证据.pdf") == "证据.pdf"


@pytest.mark.skipif(imports.docx_converter()[0] != "libreoffice", reason="No local LibreOffice for offline conversion integration test")
def test_local_libreoffice_converts_generated_docx_table(tmp_path):
    from docx import Document

    original = tmp_path / "original.docx"
    document = Document()
    document.add_paragraph("Fictional archive conversion test")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Case"
    table.cell(0, 1).text = "Fictional case 001"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Testing only"
    document.save(original)
    contents = original.read_bytes()
    target = imports.normalize_import(original, "fictional.docx")
    assert original.read_bytes() == contents
    with fitz.open(target) as pdf:
        assert pdf.page_count >= 1
        text = "".join(page.get_text() for page in pdf)
        assert "Fictional case 001" in text
        assert "Testing only" in text


@pytest.mark.skipif(imports.docx_converter()[0] != "libreoffice", reason="No local LibreOffice")
def test_real_mixed_demo_docx_converts_from_deep_uuid_import_path(tmp_path, monkeypatch):
    """Real regression: long LO profile paths crashed Windows with 0xC0000409."""
    import app.config as config
    from app.services.demo_materials import create_demo_inputs

    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    source = next(path for path in create_demo_inputs(tmp_path / "demo", "mixed") if path.suffix == ".docx")
    directory = tmp_path / "orgs" / ("a" * 32) / "cases" / ("b" * 32) / "imports" / ("c" * 32)
    directory.mkdir(parents=True)
    original = directory / "original.docx"
    original.write_bytes(source.read_bytes())
    contents = original.read_bytes()
    target = imports.normalize_import(original, source.name)
    assert original.read_bytes() == contents
    with fitz.open(target) as pdf:
        assert pdf.page_count >= 1
        assert "星河演示贸易有限公司" in "".join(page.get_text() for page in pdf)
