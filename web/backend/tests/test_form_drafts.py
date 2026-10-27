"""Only synthetic snapshots and PDFs; no browser/model/network required."""
from pathlib import Path
from types import SimpleNamespace

from docx import Document
import pytest

from app.models import TaskStatus
from app.services import archive_service, browser_pdf_service
from app.services.analysis_snapshot import save_snapshot, load_snapshot_data
from app.services.system_docx_service import render_form_docx
from app.services.task_workflow import analysis_issues
from tests.test_task_workflow import fixture_analysis
from tests.test_file_import_api import upload_context
from app.core.v4_bridge import init_v4_path

init_v4_path()


def evidence(path):
    values = {"委托人": "虚构甲公司", "案由": "虚构合同争议", "地址": "虚构路1号", "委托人电话": "00000000000"}
    return {**values, "_provenance": {key: [{"method": "local_agent", "value": value,
        "quote": key + "：" + value, "source_path": str(path), "page": 1}] for key, value in values.items()}}


def prepared(tmp_path):
    source = str(tmp_path / "synthetic.pdf")
    analysis = fixture_analysis(source, evidence(source))
    analysis.doc_spans[0].doc_type = "judgment"
    analysis.doc_spans[0].title = "虚构判决书"
    manifest = [{"file_id": "synthetic", "filename": "synthetic.pdf", "path": source, "pages": 1}]
    return analysis, manifest


def test_classification_inventory_is_not_a_substitute_for_missing_business_facts(tmp_path):
    analysis, manifest = prepared(tmp_path)
    analysis.fields = {}
    archive_service.refresh_form_drafts(analysis, {}, manifest)
    assert analysis.fields["法院文件清单"]
    issues = analysis_issues(analysis, manifest)
    assert any(item["kind"] == "empty_fields" and item["blocking"] for item in issues)
    summaries = [item for item in issues if item["kind"] == "form_draft"]
    assert len(summaries) == 1 and summaries[0]["blocking"] is False
    assert summaries[0]["detail"]["generated_count"] == 1
    assert "已送达" in summaries[0]["message"] and "不证明" in summaries[0]["message"]


@pytest.mark.asyncio
async def test_web_field_save_recomputes_drafts_and_preserves_explicit_blank(tmp_path, monkeypatch):
    analysis, manifest = prepared(tmp_path)
    archive_service.refresh_form_drafts(analysis, {}, manifest)
    work = tmp_path / "work"
    save_snapshot(str(work), analysis, base_name="synthetic")
    task = SimpleNamespace(id=101, fields=analysis.fields, manual_fields={}, source_manifest=manifest,
        issues=analysis_issues(analysis, manifest), revision=1, output_pdf="old.pdf",
        artifacts={"valid": True}, confirmation={"confirmed": True})
    monkeypatch.setattr(archive_service, "_work_dir_for", lambda *args: str(work))
    first = await archive_service.update_task_fields(task, "synthetic-org", {"委托人": "人工核对的虚构乙"}, actor_id="tester")
    assert "人工核对的虚构乙" in first["案情简介"]
    # The fixed contact cell may omit the client name repeated in other forms.
    assert "虚构路1号" in first["委托人联系地址及电话"] and "00000000000" in first["委托人联系地址及电话"]
    assert first["委托人"] == "人工核对的虚构乙"
    assert task.confirmation is None and task.artifacts is None and task.output_pdf == ""
    assert task.status == TaskStatus.awaiting_review and task.revision == 2
    await archive_service.update_task_fields(task, "synthetic-org", {"案情简介": ""}, actor_id="tester")
    final = await archive_service.update_task_fields(task, "synthetic-org", {"案由": "人工核对的咨询事项"}, actor_id="tester")
    assert final["案情简介"] == ""
    assert final["_form_drafts"]["fields"]["案情简介"]["status"] == "human_modified"
    assert load_snapshot_data(str(work))["fields"] == final == task.fields
    assert not any(key in final for key in ("主任审批意见", "已收业务费", "送达日期", "委托人对服务质量意见"))


def test_original_word_preserves_generated_drafts_and_all_eighteen_inventory_slots(tmp_path):
    import re
    from pdf_doc_locator import DocumentUnit
    analysis, manifest = prepared(tmp_path)
    manifest[0]["pages"] = 12
    analysis.doc_spans = [DocumentUnit(i, "judgment", i, i, "虚构判决书%02d" % (i + 1),
        source_path=manifest[0]["path"], confidence=1) for i in range(12)]
    fields = archive_service.refresh_form_drafts(analysis, {}, manifest)
    for name, field in (("立案审批表", "案情简介"), ("送达材料清单", "法院文件清单"), ("质量监督卡", "委托人联系地址及电话")):
        path = render_form_docx(name, fields, "虚构律所", str(tmp_path / (name + ".docx")))
        document = Document(path)
        text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        html = browser_pdf_service.render_form_html(name, fields, "虚构律所")
        for line in fields[field].splitlines():
            # Source-page annotations stay in the evidence/field value; short
            # document titles occupy the original form's 18 separate rows.
            shown = re.sub(r"（原材料第[0-9—\-]+页）$", "", line) if name == "送达材料清单" else line
            assert shown in text and line in html
    delivery = Document(tmp_path / "送达材料清单.docx")
    assert len(delivery.tables[0].rows) == 19
    for index in range(12):
        assert "虚构判决书%02d" % (index + 1) in delivery.tables[0].cell(index + 1, 1).text


def test_workflow_conflict_prevents_draft_even_when_old_provenance_still_matches(tmp_path):
    analysis, manifest = prepared(tmp_path)
    archive_service.refresh_form_drafts(analysis, {}, manifest)
    fields = archive_service.refresh_form_drafts(analysis, {}, manifest,
        issues=[{"kind": "field_conflict", "detail": {"field": "委托人"}}])
    assert fields["委托人"] == "虚构甲公司" and fields["案情简介"] == ""
    assert fields["_form_drafts"]["fields"]["案情简介"]["status"] == "insufficient"


def test_stale_human_input_needs_resolution_of_the_new_conflict_before_rederivation(tmp_path):
    from app.services.task_workflow import issue
    analysis, manifest = prepared(tmp_path)
    manual = {"委托人": analysis.fields["委托人"]}
    conflict = issue("field_conflict", "虚构补件冲突", detail={"field":"委托人"})
    archive_service.refresh_form_drafts(analysis, manual, manifest)
    fields = archive_service.refresh_form_drafts(analysis, manual, manifest, [conflict])
    assert not fields["案情简介"] and fields["委托人"] == manual["委托人"]
    fields["_issue_decisions"] = {conflict["id"]:{"action":"acknowledged", "reason":"已核对虚构补件"}}
    resolved = archive_service.refresh_form_drafts(analysis, manual, manifest, [conflict])
    assert not resolved['案情简介']  # Legacy acknowledgement does not choose the factual value.
    assert resolved['委托人']==manual['委托人']


def test_layout_capacity_issue_does_not_claim_available_material_is_missing(tmp_path):
    analysis, manifest = prepared(tmp_path)
    analysis.fields["_form_drafts"] = {"fields":{"委托人联系地址及电话":{
        "status":"insufficient", "reason_codes":["draft_too_long_for_original_cell"]}}}
    summary = next(item for item in analysis_issues(analysis,manifest) if item["kind"] == "form_draft")
    assert "原表空间有限" in summary["message"]
    assert "委托人联系地址及电话" in summary["detail"]["layout_review_fields"]
    assert "委托人联系地址及电话" not in summary["detail"]["insufficient_fields"]


@pytest.mark.parametrize("missing_tail", [False, True])
def test_long_inventory_pdf_cannot_silently_drop_tail(tmp_path, missing_tail):
    import fitz
    lines = [f"SyntheticDocument-{i:02d}" for i in range(12)]
    path = tmp_path / "synthetic-inventory.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "\n".join(lines[:-1] if missing_tail else lines))
        document.save(path)
    if missing_tail:
        with pytest.raises(browser_pdf_service.BrowserPdfError, match="BROWSER_INVENTORY_INCOMPLETE") as error:
            browser_pdf_service._validate_long_inventory(path, "送达材料清单", {"法院文件清单": "\n".join(lines)})
        assert "SyntheticDocument" not in str(error.value) and str(path) not in str(error.value)
    else:
        browser_pdf_service._validate_long_inventory(path, "送达材料清单", {"法院文件清单": "\n".join(lines)})


@pytest.mark.asyncio
@pytest.mark.parametrize("manual_inventory", [False, True])
async def test_reclassifying_court_document_refreshes_inventory_without_overwriting_human(upload_context, monkeypatch, manual_inventory):
    from tests.test_document_classification import fixture_task, edit
    from app.models import ArchiveTask
    from app.services.analysis_snapshot import load_analysis
    ctx = upload_context
    task_id, work, manifest = await fixture_task(ctx, monkeypatch)
    analysis, snapshot = load_analysis(work)
    manual = {"合同号": "人工保留"}
    if manual_inventory:
        manual["法院文件清单"] = "人工核对的虚构清单"
    archive_service.refresh_form_drafts(analysis, manual, manifest)
    async with ctx.sessions() as db:
        task = await db.get(ArchiveTask, task_id)
        task.fields = analysis.fields
        task.manual_fields = manual
        await db.commit()
        save_snapshot(work, analysis, base_name="虚构重分类", metadata={"source_version": task.source_version, "revision": 1})
    rows = (await ctx.client.get(f"/api/tasks/{task_id}/documents")).json()["documents"]
    response = await edit(ctx, task_id, rows[1]["unit_key"], doc_type="evidence", catalog_seq=7,
        reason="虚构反例：核对后确认是证据附件")
    assert response.status_code == 200, response.text
    task_body = response.json()["task"]
    assert task_body["revision"] == 2
    expected = "人工核对的虚构清单" if manual_inventory else ""
    assert task_body["fields"]["法院文件清单"] == expected
    assert load_snapshot_data(work)["fields"]["法院文件清单"] == expected
    status = task_body["fields"]["_form_drafts"]["fields"]["法院文件清单"]["status"]
    assert status == ("human_modified" if manual_inventory else "insufficient")
    assert any(item["kind"] == "ocr" for item in task_body["issues"])
