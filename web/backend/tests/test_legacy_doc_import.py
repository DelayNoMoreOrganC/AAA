"""Static legacy DOC allowlist and actual local conversion of fictional data."""
import shutil
import struct
import subprocess
from pathlib import Path
import sys

import fitz
import olefile
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services import file_import_service as imports


@pytest.fixture(scope="module")
def fictional_doc(tmp_path_factory):
    executable = imports.legacy_doc_converter()
    if not executable:
        pytest.skip("No guarded local LibreOffice legacy conversion")
    from docx import Document
    folder = tmp_path_factory.mktemp("fictional-doc")
    source = folder / "fictional.docx"
    doc = Document()
    doc.add_paragraph("Fictional legacy archive")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Fictional case"
    table.cell(0, 1).text = "TEST 2026 001"
    table.cell(1, 0).text = "Purpose"
    table.cell(1, 1).text = "Generated test data only"
    doc.save(source)
    profile = folder / "creation-profile"
    imports._secure_office_profile(profile)
    result = subprocess.run([executable, "-env:UserInstallation=" + profile.as_uri(),
                             "--headless", "--nologo", "--norestore", "--convert-to",
                             "doc:MS Word 97", "--outdir", str(folder), str(source)],
                            capture_output=True, timeout=120, check=False,
                            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0)
    assert result.returncode == 0
    target = folder / "fictional.doc"
    assert target.is_file()
    return target


def change_word(source, target, mutate):
    shutil.copyfile(source, target)
    with olefile.OleFileIO(target, write_mode=True) as ole:
        data = bytearray(ole.openstream("WordDocument").read())
        mutate(data)
        ole.write_stream("WordDocument", bytes(data))


def test_legacy_requires_ole_parser_and_guarded_libreoffice_not_word(monkeypatch):
    monkeypatch.setattr(imports, "_libreoffice_path", lambda: "")
    monkeypatch.setattr(imports, "docx_converter", lambda: ("word", sys.executable))
    assert ".docx" in imports.import_capabilities(100)["supported_extensions"]
    assert ".doc" not in imports.import_capabilities(100)["supported_extensions"]
    with pytest.raises(imports.ImportFailure, match="安全转换"):
        imports.check_extension("test.doc")


def test_renamed_rtf_never_reaches_office(tmp_path, monkeypatch):
    source = tmp_path / "test.doc"
    source.write_bytes(b"{\\rtf1 fictional}")
    monkeypatch.setattr(imports.subprocess, "run", lambda *a, **kw: pytest.fail("Office must not open unchecked DOC"))
    with pytest.raises(imports.ImportFailure, match="改名"):
        imports.normalize_import(source, "test.doc")


@pytest.mark.parametrize("pair", [16, 17, 18, 19, 20, 48, 57, 59])
def test_all_document_part_fields_are_rejected_before_conversion(fictional_doc, tmp_path, monkeypatch, pair):
    source = tmp_path / "field.doc"
    change_word(fictional_doc, source, lambda data: struct.pack_into("<II", data, 154 + pair * 8, 0, 1))
    monkeypatch.setattr(imports.subprocess, "run", lambda *a, **kw: pytest.fail("Office must not see DOC fields"))
    with pytest.raises(imports.ImportFailure, match="字段域"):
        imports.normalize_import(source, "field.doc")


@pytest.mark.parametrize("pair", [24, 26, 40, 45, 50, 62, 72, 86, 99])
def test_external_references_objects_and_unknown_extensions_fail_closed(fictional_doc, tmp_path, pair):
    source = tmp_path / "complex.doc"
    change_word(fictional_doc, source, lambda data: struct.pack_into("<II", data, 154 + pair * 8, 0, 1))
    with pytest.raises(imports.ImportFailure, match="外部引用"):
        imports._validate_doc(source)


def test_encryption_rejected(fictional_doc, tmp_path):
    source = tmp_path / "encrypted.doc"
    def encrypt_flag(data):
        struct.pack_into("<H", data, 10, struct.unpack_from("<H", data, 10)[0] | 0x100)
    change_word(fictional_doc, source, encrypt_flag)
    with pytest.raises(imports.ImportFailure, match="加密"):
        imports._validate_doc(source)


def test_unindexed_field_control_is_still_rejected(fictional_doc, tmp_path):
    source = tmp_path / "unindexed.doc"
    # The genuine generated fixture's first text piece starts at offset 2048.
    # Obtain its actual offset rather than relying on that normal layout.
    with olefile.OleFileIO(fictional_doc) as ole:
        word = ole.openstream("WordDocument").read()
        table = ole.openstream("1Table" if struct.unpack_from("<H", word, 10)[0] & 0x200 else "0Table").read()
        offset, length = struct.unpack_from("<II", word, 154 + 33 * 8)
        clx = table[offset:offset + length]
        count = (struct.unpack_from("<I", clx, 1)[0] - 4) // 12
        fc = struct.unpack_from("<I", clx, 5 + (count + 1) * 4 + 2)[0]
        text_offset = (fc & 0x3FFFFFFF) // (2 if fc & 0x40000000 else 1)
    change_word(fictional_doc, source, lambda data: data.__setitem__(text_offset, 0x13))
    with pytest.raises(imports.ImportFailure, match="动态字段"):
        imports._validate_doc(source)


def test_macro_named_stream_is_rejected(fictional_doc, tmp_path):
    source = tmp_path / "macro.doc"
    data = fictional_doc.read_bytes()
    before, after = "1Table".encode("utf-16le"), "Macros".encode("utf-16le")
    assert data.count(before) == 1
    source.write_bytes(data.replace(before, after))
    with pytest.raises(imports.ImportFailure, match="宏"):
        imports._validate_doc(source)


def test_associated_external_template_rejected(fictional_doc, tmp_path):
    source = tmp_path / "template.doc"
    shutil.copyfile(fictional_doc, source)
    strings = [""] * 18
    strings[1] = "https://invalid.test/template.dot"
    record = struct.pack("<HHH", 0xFFFF, 18, 0) + b"".join(struct.pack("<H", len(value)) + value.encode("utf-16le") for value in strings)
    with olefile.OleFileIO(source, write_mode=True) as ole:
        word = bytearray(ole.openstream("WordDocument").read())
        table = bytearray(ole.openstream("1Table").read())
        table[:len(record)] = record
        struct.pack_into("<II", word, 154 + 32 * 8, 0, len(record))
        ole.write_stream("WordDocument", bytes(word))
        ole.write_stream("1Table", bytes(table))
    with pytest.raises(imports.ImportFailure, match="外部模板"):
        imports._validate_doc(source)


def test_out_of_bounds_piece_is_not_repaired_or_opened(fictional_doc, tmp_path):
    source = tmp_path / "bad-table.doc"
    change_word(fictional_doc, source, lambda data: struct.pack_into("<II", data, 154 + 33 * 8, 0x7FFFFFFF, 21))
    with pytest.raises(imports.ImportFailure, match="结构无法完整核验"):
        imports._validate_doc(source)


def test_guarded_profile_disables_macros_active_content_and_link_updates(tmp_path):
    from xml.etree import ElementTree
    imports._secure_office_profile(tmp_path / "profile")
    tree = ElementTree.parse(tmp_path / "profile/user/registrymodifications.xcu")
    values = {node.attrib["{http://openoffice.org/2001/registry}name"]: node.findtext("value") for node in tree.iter("prop")}
    assert values["DisableMacrosExecution"] == "true"
    assert values["DisableActiveContent"] == "true"
    assert values["BlockUntrustedRefererLinks"] == "true"
    assert values["SecureURL"] == ""
    assert values["Link"] == "2"
    assert values["Field"] == "false"


def test_actual_fictional_doc_keeps_original_and_converts_text_and_table(fictional_doc, tmp_path):
    original = tmp_path / "original.doc"
    shutil.copyfile(fictional_doc, original)
    before = original.read_bytes()
    target = imports.normalize_import(original, "fictional.doc")
    assert original.read_bytes() == before
    with fitz.open(target) as pdf:
        assert pdf.page_count >= 1
        text = "".join(page.get_text() for page in pdf)
        assert "Fictional legacy archive" in text
        assert "TEST 2026 001" in text
        assert "Generated test data only" in text
