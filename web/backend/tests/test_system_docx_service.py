from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from app.services.analysis_snapshot import SYSTEM_TEMPLATE_NAMES
from app.services.browser_pdf_service import FORM_LAYOUTS as PDF_LAYOUTS, render_form_html
from app.services.form_layouts import FORM_LAYOUTS
from app.services.system_docx_service import generate_system_docx, render_form_docx
from app.services.original_template_service import OriginalTemplateError, template_metadata, TEMPLATES, FILES


def _all_text(document):
    return "".join(document.element.itertext())


def test_all_five_documents_use_confirmed_values_and_shared_layout(tmp_path):
    fields = {"案件类别": "行政", "委托人": "虚构测试甲方", "承办律师": "虚构律师",
              "法院收案号": "虚构案号", "结案小结": "尚未履行和解协议。",
              "委托人对服务质量意见": "希望改善沟通。", "法院文件清单": "虚构判决书\n虚构裁定书",
              "_preview_custom_values": {"立案审批表": {"p0-b0-r6-c0": "已经人工核对。"}}}
    paths = generate_system_docx(fields, "虚构律所", tmp_path, log=lambda *_a: None)
    assert set(paths) == set(SYSTEM_TEMPLATE_NAMES)
    assert PDF_LAYOUTS is FORM_LAYOUTS
    for name, path in paths.items():
        document = Document(path)
        assert document.tables
        assert document.settings.element.find(qn("w:documentProtection")).get(qn("w:edit")) == "readOnly"
        assert len(document.element.findall(".//" + qn("w:permStart"))) == len(document.element.findall(".//" + qn("w:permEnd"))) > 0
    report = _all_text(Document(paths["结案报告表"]))
    assert "尚未履行和解协议。" in report and "希望改善沟通。" in report
    assert "满意" not in report
    assert "已经人工核对。" in _all_text(Document(paths["立案审批表"]))
    assert "已经人工核对。" in render_form_html("立案审批表", fields, "虚构律所")
    delivery = _all_text(Document(paths["送达材料清单"]))
    assert "虚构判决书" in delivery and "虚构裁定书" in delivery


def test_fixed_labels_have_no_edit_permission_and_ignore_custom_injection(tmp_path):
    path = render_form_docx("立案审批表", {
        "委托人": "虚构甲",
        "_preview_custom_values": {"立案审批表": {"p0-b0-r0-c0": "伪造固定标签"}},
        "_preview_styles": {"立案审批表": {"field:委托人": {"fontSize": "9pt", "fontWeight": "bold"}}},
    }, "虚构律所", str(tmp_path / "form.docx"))
    document = Document(path)
    assert "伪造固定标签" not in _all_text(document)
    fixed = document.tables[0].cell(0, 0)
    assert fixed.text == "案件类别"
    assert not fixed._tc.findall(".//" + qn("w:permStart"))
    value = document.tables[0].cell(1, 1)
    text_run = next(run for run in value.paragraphs[0].runs if "虚构甲" in run.text)
    assert text_run.font.size.pt == 12 and text_run.bold is not True
    assert text_run.font.name == "宋体"


def test_large_value_is_rejected_without_truncation_or_growing_original_rows(tmp_path):
    value = "虚构长内容。" * 200
    fields = {"结案小结": value}
    with pytest.raises(OriginalTemplateError) as error:
        render_form_docx("结案报告表", fields, "虚构律所", str(tmp_path / "long.docx"))
    assert error.value.code == "overflow" and error.value.fields == ["结案小结"]
    assert value not in str(error.value) and fields["结案小结"] == value
    assert not (tmp_path / "long.docx").exists()


def test_blank_customer_feedback_stays_blank(tmp_path):
    path = render_form_docx("结案报告表", {}, "虚构律所", str(tmp_path / "blank.docx"))
    assert "满意" not in _all_text(Document(path))


def test_nonlitigation_cover_title_matches_html_and_has_no_title_rule(tmp_path):
    fields = {"案件类别": "法律顾问"}
    path = render_form_docx("档案卷宗", fields, "虚构律所", str(tmp_path / "counsel.docx"))
    document = Document(path)
    assert "非诉讼类" in _all_text(document)
    assert "非诉讼类" in render_form_html("档案卷宗", fields, "虚构律所")
    assert not document.styles.element.xpath(".//w:pBdr")


def test_blank_category_uses_original_default_cover_without_fabricating_category(tmp_path):
    path = render_form_docx("档案卷宗", {}, "虚构律所", str(tmp_path / "unknown.docx"))
    document = Document(path)
    assert document.tables[0].cell(0, 1).text.replace("\u200b", "") == ""
    assert "民事" not in _all_text(document)
    assert template_metadata("档案卷宗", {})["template_key"] == "档案卷宗"


def test_quality_card_preserves_original_notice_with_historical_metadata(tmp_path):
    path = render_form_docx("质量监督卡", {}, "虚构律所", str(tmp_path / "quality.docx"))
    original = Document(TEMPLATES / "word-converted" / (FILES["质量监督卡"] + ".docx"))
    document = Document(path)
    assert [p.text for p in document.paragraphs] == [p.text for p in original.paragraphs]
    assert template_metadata("质量监督卡")["historical_template"] is True
